"""Read-only VA STOCK viewer for Render; inventory comes from a secret file."""
import json
import os
import unicodedata
from io import BytesIO
from pathlib import Path
from typing import Literal

from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent
INVENTORY_PATH = Path(os.environ.get("INVENTORY_PATH", "/etc/secrets/inventory.json"))
app = FastAPI(title="VA STOCK", docs_url=None, redoc_url=None, openapi_url=None)
app.mount("/static", StaticFiles(directory=ROOT / "static"), name="static")


@app.middleware("http")
async def read_only_viewer(request, call_next):
    headers = {"Cache-Control": "no-store", "X-Robots-Tag": "noindex, nofollow"}
    if request.method not in ("GET", "HEAD"):
        return Response(status_code=403, headers=headers)
    response = await call_next(request)
    response.headers.update(headers)
    return response


@app.get("/healthz")
def health():
    return {"status": "ok"}


@app.get("/")
def home():
    return FileResponse(ROOT / "templates" / "index.html")


@app.get("/api/inventory")
def inventory():
    try:
        body = json.loads(INVENTORY_PATH.read_text(encoding="utf-8"))
        if not isinstance(body, dict) or not isinstance(body.get("data"), list):
            raise ValueError("Invalid inventory snapshot")
        report = body.get("report")
        # Match older snapshots to the known report metadata by their exact timestamp.
        if not report and (ROOT / "report_context.json").exists():
            context = json.loads((ROOT / "report_context.json").read_text(encoding="utf-8"))
            if context.get("updated_at") == body.get("updated_at"):
                report = context.get("report")
        return {"success": True, "count": len(body["data"]), "report": report,
                "updated_at": body.get("updated_at"), "data": body["data"]}
    except (OSError, ValueError) as exc:
        raise HTTPException(status_code=503, detail="Chưa có dữ liệu tồn kho hợp lệ. Hãy thêm Secret File inventory.json trên Render.") from exc


def folded(value):
    text = unicodedata.normalize("NFD", str(value).casefold())
    return "".join(char for char in text if unicodedata.category(char) != "Mn").replace("đ", "d")


@app.get("/api/export.xlsx")
def export_inventory(q: str = "", stock: Literal["all", "available", "empty", "negative"] = "all",
                     sort: Literal["ma_hang", "ten_hang", "ton_kho"] = "ma_hang", descending: bool = False):
    query = folded(q.strip())
    rows = [row for row in inventory()["data"] if query in folded(row["ma_hang"]) or query in folded(row["ten_hang"])]
    if stock != "all":
        rows = [row for row in rows if (stock == "available" and row["ton_kho"] > 0)
                or (stock == "empty" and row["ton_kho"] <= 0) or (stock == "negative" and row["ton_kho"] < 0)]
    rows.sort(key=lambda row: row[sort] if sort == "ton_kho" else folded(row[sort]), reverse=descending)
    book = Workbook()
    sheet = book.active
    sheet.title = "Tồn kho"
    sheet.append(["STT", "Mã hàng", "Tên hàng", "Tồn kho"])
    for index, row in enumerate(rows, start=1):
        sheet.append([index, str(row["ma_hang"]), str(row["ten_hang"]), row["ton_kho"]])
        # Item labels are text even when they start with an Excel formula character.
        for cell in sheet[sheet.max_row][1:3]:
            cell.data_type = "s"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="173B52")
        cell.font = Font(name="Segoe UI", color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    for column, width in zip("ABCD", (8, 20, 85, 18)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 26
    output = BytesIO()
    book.save(output)
    return Response(content=output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="TonKho.xlsx"'})
