"""Browser checks against a running VA STOCK viewer with inventory data."""
import os
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright


def check_ui():
    base = os.environ.get("VA_STOCK_TEST_URL", "http://127.0.0.1:8012")
    artifacts = Path(os.environ.get("VA_STOCK_ARTIFACTS", "../data/render/ui-check"))
    artifacts.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 1920, "height": 1030}, accept_downloads=True)
        errors = []
        page.on("pageerror", lambda error: errors.append(str(error)))
        page.goto(base)
        page.wait_for_load_state("networkidle")
        expect(page.locator("#rows tr").first).to_be_visible()
        total = int(page.locator("#count").inner_text().replace(".", ""))
        available = int(page.locator("#available-count").inner_text().replace(".", ""))
        empty = int(page.locator("#empty-count").inner_text().replace(".", ""))
        assert total == available + empty
        expect(page.locator("#rows tr")).to_have_count(total)
        page.screenshot(path=str(artifacts / "desktop.png"), full_page=True)

        page.locator('[data-stock="available"]').click()
        expect(page.locator("#rows tr")).to_have_count(available)
        page.locator('[data-stock="empty"]').click()
        expect(page.locator("#rows tr")).to_have_count(empty)
        page.locator("#clear-filters").click()
        page.locator("#query").fill("binh thuy tinh")
        expect(page.locator("#rows tr").first).to_contain_text("Bình thủy tinh")
        found = page.locator("#rows tr").count()
        assert 0 < found < total
        with page.expect_download() as download:
            page.locator("#export").click()
        workbook = load_workbook(BytesIO(Path(download.value.path()).read_bytes()))
        assert workbook.active.max_row == found + 1

        page.locator("#query").fill("no-such-product-938493")
        expect(page.locator("#empty")).to_be_visible()
        expect(page.locator("#export")).to_be_disabled()
        page.locator("#clear-filters").click()
        page.locator("#more-filters").click()
        expect(page.get_by_role("dialog")).to_be_visible()
        page.locator("#stock-filter").select_option("available")
        page.get_by_role("button", name="Áp dụng", exact=True).click()
        expect(page.locator("#rows tr")).to_have_count(available)
        page.locator("#clear-filters").click()
        page.locator('[data-sort="ton_kho"]').click()
        quantities = [text.strip() for text in page.locator("#rows tr td:last-child").all_text_contents()]
        assert quantities[0] == "Hết hàng" or quantities[0].startswith("⚠")
        page.locator('[data-sort="ton_kho"]').click()
        assert "Hết hàng" not in page.locator("#rows tr").first.inner_text()
        page.locator('[data-sort="ma_hang"]').click()
        page.locator("#refresh").click()
        expect(page.locator("#status")).to_contain_text("Đã tải bản tồn kho")

        page.locator('[data-view="settings"]').click()
        expect(page.locator("#settings")).to_be_visible()
        page.locator("#theme").select_option("dark")
        page.locator("#refresh-interval").select_option("300")
        page.reload()
        page.wait_for_load_state("networkidle")
        expect(page.locator("html")).to_have_attribute("data-theme", "dark")
        page.locator('[data-view="settings"]').click()
        expect(page.locator("#refresh-interval")).to_have_value("300")
        page.locator("#theme").select_option("light")
        page.locator('[data-view="inventory"]').click()
        expect(page.locator("#query")).to_be_focused()

        for width in (390, 360):
            page.set_viewport_size({"width": width, "height": 844})
            assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth"), f"Page overflow at {width}px"
            expect(page.locator("#more-filters")).to_be_visible()
            quantity = page.locator("#rows tr td:last-child").first.bounding_box()
            assert quantity["x"] + quantity["width"] <= width, "Stock column is hidden on mobile"
            if width == 390:
                page.screenshot(path=str(artifacts / "mobile.png"), full_page=True)
        assert not errors, errors
        print(f"PASS: desktop/mobile layout, {total} rows, counts, accent-insensitive search, stock filters, sorting, Excel, settings and refresh")
        browser.close()


if __name__ == "__main__":
    check_ui()
